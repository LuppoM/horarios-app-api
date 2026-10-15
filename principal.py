from datetime import date
from io import BytesIO
import os
from typing import Optional

from fastapi import Body, FastAPI, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from abmmodelo import abmmodmoimodelo
from dtprogramadamodelo import ModeloDTProgramadaWeb
from Hchocomodelo import ModeloHChocoWeb
from hchocofindemodelo import ModeloHChocoFindeWeb
from hgolofindemodelo import ModeloHGoloFindeWeb
from hgolomodelo import ModeloHGoloWeb
from mecanicosmodelo import ModeloMecanicosWeb
from modmoimodelo import ModMoiModel
from principalmodelo import modeloprincipal
import sectoresmodelo as modelo_sectores
import tipohorariosmodelo as modelo_tipohorarios

modelo_dtprog = ModeloDTProgramadaWeb()
app = FastAPI(title="ARMADO DE HORARIOS")

# Instanciamos los modelos generales
modmoi_modelo = ModMoiModel()
abm_modelo = abmmodmoimodelo()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
templates = Jinja2Templates(directory=BASE_DIR)

# Montar archivos estáticos
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")
app.mount("/archivos", StaticFiles(directory=os.path.join(BASE_DIR, "archivos")), name="archivos")


def semana_objetivo_por_defecto() -> str:
    """Obtiene la semana activa actual según la regla operativa."""
    return str(
        modelo_dtprog.obtener_numero_semana_personalizado(date.today())
    )


# ==========================================
# 1. VISTAS HTML PRINCIPALES (GENERAL)
# ==========================================

@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    modelo_base = modeloprincipal()
    try:
        # Limpieza global de semanas anteriores en Supabase (< semana actual)
        modelo_base.eliminar_datos_semanas_anteriores()
    except Exception as e:
        print(f"Error al ejecutar limpieza automatica en el Inicio: {e}")
    finally:
        if hasattr(modelo_base, 'cerrar_conexion'):
            modelo_base.cerrar_conexion()

    return templates.TemplateResponse(request, "Principal.html", {})


@app.get("/modmoi", response_class=HTMLResponse)
def vista_modmoi(request: Request):
    return templates.TemplateResponse(request, "modmoi.html", {})


@app.get("/sectores", response_class=HTMLResponse)
def vista_sectores(request: Request):
    return templates.TemplateResponse(request, "sectores.html", {})


@app.get("/tipohorarios", response_class=HTMLResponse)
@app.get("/tiposhorarios", response_class=HTMLResponse)
def vista_tipos_horario(request: Request):
    return templates.TemplateResponse(request, "tipohorarios.html", {})


@app.get("/hchoco", response_class=HTMLResponse)
def vista_hchoco(request: Request):
    return templates.TemplateResponse(request, "hchoco.html", {})


@app.get("/dtprogramada", response_class=HTMLResponse)
def vista_dotacion_programada(request: Request):
    return templates.TemplateResponse(request, "dtprogramada.html", {})


@app.get("/hchocofinde", response_class=HTMLResponse)
def vista_hchocofinde(request: Request):
    return templates.TemplateResponse(request, "hchocofinde.html", {})


@app.get("/hgolo", response_class=HTMLResponse)
def vista_hgolo(request: Request):
    return templates.TemplateResponse(request, "hgolo.html", {})


@app.get("/hgolofinde", response_class=HTMLResponse)
def vista_hgolo_finde(request: Request):
    return templates.TemplateResponse(request, "hgolofinde.html", {})


@app.get("/mecanicos", response_class=HTMLResponse)
def vista_mecanicos(request: Request):
    return templates.TemplateResponse(request, "mecanicos.html", {})
# ==========================================
# 1.1. ESTADO PARA LA VISTA PRINCIPAL (INICIO)
# ==========================================

@app.get("/api/dtprogramada/datos")
def api_obtener_dtprogramada(semana: Optional[str] = None):
    if not semana or semana.lower() == "none" or semana.strip() == "":
        semana = str(modelo_dtprog.obtener_numero_semana_personalizado())

    # Fuerza una conexión limpia en esta llamada
    conn = None
    cursor = None
    es_finalizado = False
    modificado_por = ""
    fecha_modificacion = ""
    estado_str = "PENDIENTE"

    try:
        conn = conexion.obtener_conexion()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT estado, modificado_por, fecha_modificacion FROM estados_programacion WHERE semana::text = %s;", 
            (str(semana).strip(),)
        )
        row = cursor.fetchone()
        if row:
            estado_str = str(row[0]).strip().upper() if row[0] else "PENDIENTE"
            es_finalizado = (estado_str == "FINALIZADO")
            modificado_por = str(row[1]) if row[1] else ""
            fecha_modificacion = str(row[2]) if row[2] else ""
    except Exception as err:
        print(f"Error real leyendo BD: {err}")
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()

    try:
        sectores_choco = modelo_dtprog.obtener_sectores_rotativos_por_planta("Chocolate")
        sectores_golo = modelo_dtprog.obtener_sectores_rotativos_por_planta("Golosinas")
        turnos = modelo_dtprog.obtener_turnos_por_tipo("ROTATIVO")
        dotaciones = modelo_dtprog.obtener_dotaciones_semana(semana)

        return {
            "semana": str(semana),
            "estado": estado_str,
            "finalizado": es_finalizado,
            "modificado_por": modificado_por,
            "fecha_modificacion": fecha_modificacion,
            "turnos": turnos,
            "sectores_chocolate": sectores_choco,
            "sectores_golosinas": sectores_golo,
            "dotaciones": dotaciones
        }
    except Exception as e:
        return {"error": str(e)}
# ==========================================
# 2. MÓDULO: MANO DE OBRA (MOD/MOI)
# ==========================================

@app.get("/api/modmoi/datos")
def api_obtener_modmoi(
    vista: str = "Semana", columna: str = None, filtro: str = None
):
    if vista == "Mecanicos":
        columnas, datos = modmoi_modelo.obtener_datos_mecanicos(columna, filtro)
    elif vista == "Fin de Semana":
        columnas, datos = modmoi_modelo.obtener_datos_finde(columna, filtro)
    else:
        columnas, datos = modmoi_modelo.obtener_datos(columna, filtro)

    mod_count, moi_count = modmoi_modelo.contar_mano_de_obra(vista=vista)

    return {
        "columnas": columnas,
        "datos": datos,
        "mod": mod_count,
        "moi": moi_count,
    }


@app.get("/api/modmoi/exportar-excel")
def api_exportar_excel(vista: str = "Semana"):
    semana_actual = modmoi_modelo.obtener_numero_semana_personalizado(date.today())

    if vista == "Mecanicos":
        columnas, datos = modmoi_modelo.obtener_datos_mecanicos(semana_actual=semana_actual)
        nombre_hoja = "Horarios Mecanicos"
        nombre_archivo = f"Consolidado Mecanicos semana {semana_actual}.xlsx"
    elif vista == "Fin de Semana":
        columnas, datos = modmoi_modelo.obtener_datos_finde(semana_actual=semana_actual)
        nombre_hoja = "Horarios Finde"
        nombre_archivo = f"Consolidado Finde semana {semana_actual}.xlsx"
    else:
        columnas, datos = modmoi_modelo.obtener_datos(semana_actual=semana_actual)
        nombre_hoja = "Horarios Semana"
        nombre_archivo = f"Consolidado Semana {semana_actual}.xlsx"

    if not datos:
        return {"error": "No hay datos registrados para exportar."}

    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    encabezado_font = Font(bold=True, color="FFFFFF")
    encabezado_fill = PatternFill("solid", fgColor="4F81BD")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for i, col in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i, value=col)
        celda.font = encabezado_font
        celda.fill = encabezado_fill
        celda.border = thin_border
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for fila_idx, fila in enumerate(datos, start=2):
        for col_idx, valor in enumerate(fila, start=1):
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            celda.border = thin_border
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    for col_idx, col in enumerate(columnas, start=1):
        max_length = len(str(col))
        for fila in datos:
            val_celda = fila[col_idx - 1]
            texto_val = str(val_celda) if val_celda is not None else ""
            if len(texto_val) > max_length:
                max_length = len(texto_val)

        letra_col = get_column_letter(col_idx)
        ws.column_dimensions[letra_col].width = max(15, min(max_length + 2, 50))

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    headers = {"Content-Disposition": f'attachment; filename="{nombre_archivo}"'}

    return StreamingResponse(
        output,
        headers=headers,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

@app.get("/api/modmoi/obtener/{id_registro}")
def api_obtener_registro(id_registro: int):
    data = abm_modelo.obtener_por_id(id_registro)
    if not data:
        return {"error": "Registro no encontrado"}
    return data


@app.post("/api/modmoi/guardar")
def api_guardar_registro(payload: dict = Body(...)):
    modo = payload.get("modo")
    id_registro = payload.get("id")
    datos = payload.get("datos", {})

    if not datos.get("calificacion_mano_obra"):
        return {"status": "error", "message": "Seleccione Calificación de Mano de Obra"}
    if not datos.get("operario"):
        return {"status": "error", "message": "El nombre de operario es requerido"}
    if not datos.get("horario"):
        return {"status": "error", "message": "El horario es requerido"}

    try:
        if modo == "insertar":
            abm_modelo.insertar(datos)
        elif modo == "modificar":
            abm_modelo.modificar(id_registro, datos)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.delete("/api/modmoi/eliminar/{id_registro}")
def api_eliminar_registro(id_registro: int):
    try:
        abm_modelo.eliminar(id_registro)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.post("/api/modmoi/limpiar-semanas")
def api_limpiar_semanas_manual():
    try:
        modmoi_modelo.limpiar_semanas_anteriores()
        return {"status": "ok", "message": "Puestos blanqueados correctamente en Supabase."}
    except Exception as e:
        return {"status": "error", "message": str(e)}


# ==========================================
# 3. MÓDULO: SECTORES
# ==========================================

@app.get("/api/sectores/datos")
def api_obtener_sectores():
    filas = modelo_sectores.obtener_sectores()
    tipos_horario = modelo_sectores.obtener_tipos_horario()
    columnas = ["ID", "Sector", "Planta", "Tipo de Horario", "Requiere Dotación"]
    return {
        "columnas": columnas,
        "datos": filas,
        "tipos_horario": tipos_horario
    }


@app.get("/api/sectores/obtener/{id_sector}")
def api_obtener_sector(id_sector: int):
    data = modelo_sectores.obtener_sector_por_id(id_sector)
    if not data:
        return {"error": "Sector no encontrado"}
    return data


@app.post("/api/sectores/guardar")
def api_guardar_sector(payload: dict = Body(...)):
    modo = payload.get("modo")
    id_sector = payload.get("id")
    datos = payload.get("datos", {})

    try:
        if modo == "insertar":
            modelo_sectores.insertar_sector(datos)
        else:
            modelo_sectores.modificar_sector(id_sector, datos)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.delete("/api/sectores/eliminar/{id_sector}")
def api_eliminar_sector(id_sector: int):
    try:
        modelo_sectores.eliminar_sector(id_sector)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


# ==========================================
# 4. MÓDULO: TIPOS DE HORARIO
# ==========================================

@app.get("/api/tipohorarios/datos")
def api_obtener_tipos_horarios():
    filas = modelo_tipohorarios.obtener_tipos_horarios()
    columnas = ["ID", "Tipo", "Horarios / Turnos"]
    return {
        "columnas": columnas,
        "datos": filas
    }


@app.get("/api/tipohorarios/obtener/{tipo_id}")
def api_obtener_tipo_horario(tipo_id: int):
    data = modelo_tipohorarios.obtener_tipo_horario_por_id(tipo_id)
    if not data:
        return {"error": "Tipo de horario no encontrado"}
    return data


@app.post("/api/tipohorarios/guardar")
def api_guardar_tipo_horario(payload: dict = Body(...)):
    modo = payload.get("modo")
    tipo_id = payload.get("id")
    datos = payload.get("datos", {})

    if not datos.get("tipo") or not datos.get("horarios_definidos"):
        return {"status": "error", "message": "Complete el tipo y al menos un horario."}

    try:
        if modo == "insertar":
            modelo_tipohorarios.insertar_tipo_horario(datos)
        else:
            modelo_tipohorarios.modificar_tipo_horario(tipo_id, datos)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.delete("/api/tipohorarios/eliminar/{tipo_id}")
def api_eliminar_tipo_horario(tipo_id: int):
    try:
        modelo_tipohorarios.eliminar_tipo_horario(tipo_id)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


# ==========================================
# 5. MÓDULO: DOTACIÓN PROGRAMADA
# ==========================================

# ==========================================
# MÓDULO: DOTACIÓN PROGRAMADA
# ==========================================

# ==========================================
# MÓDULO: DOTACIÓN PROGRAMADA
# ==========================================

@app.get("/api/dtprogramada/datos")
def api_obtener_dtprogramada(semana: Optional[str] = None):
    """
    Obtiene los datos de la dotación y el estado de la programación.
    Si no se especifica semana, busca la última semana activa o finalizada en la BD.
    """
    try:
        # Si no envían semana, intentamos buscar la última registrada en estados_programacion
        if not semana or semana.lower() == "none" or semana.strip() == "":
            cursor = modelo_dtprog.conexion_db.cursor()
            cursor.execute("SELECT semana FROM estados_programacion ORDER BY fecha_modificacion DESC LIMIT 1;")
            row = cursor.fetchone()
            if row:
                semana = str(row[0] if not isinstance(row, dict) else row.get("semana"))
            else:
                semana = str(modelo_dtprog.obtener_numero_semana_personalizado())
            if cursor:
                cursor.close()

        sectores_choco = modelo_dtprog.obtener_sectores_rotativos_por_planta("Chocolate")
        sectores_golo = modelo_dtprog.obtener_sectores_rotativos_por_planta("Golosinas")
        turnos = modelo_dtprog.obtener_turnos_por_tipo("ROTATIVO")
        dotaciones = modelo_dtprog.obtener_dotaciones_semana(semana)

        info_estado = modelo_dtprog.obtener_estado_programacion(semana)
        es_finalizado = (info_estado.get("estado") == "FINALIZADO")

        if es_finalizado:
            mensaje = f"✅ Dotación Semana {semana} FINALIZADA por {info_estado.get('modificado_por')} ({info_estado.get('fecha_modificacion')})"
        else:
            mensaje = f"⚠️ Dotación Semana {semana} PENDIENTE de carga"

        return {
            "semana": str(semana),
            "estado": info_estado.get("estado", "PENDIENTE"),
            "finalizado": es_finalizado,
            "modificado_por": info_estado.get("modificado_por", ""),
            "fecha_modificacion": info_estado.get("fecha_modificacion", ""),
            "mensaje": mensaje,
            "turnos": turnos,
            "sectores_chocolate": sectores_choco,
            "sectores_golosinas": sectores_golo,
            "dotaciones": dotaciones
        }
    except Exception as e:
        print(f"Error en /api/dtprogramada/datos: {e}")
        return {"error": str(e)}


@app.post("/api/dtprogramada/guardar")
def api_guardar_dtprogramada(payload: dict = Body(...)):
    """
    Guarda los datos de la dotación (enviados en formato jerárquico por planta) 
    y actualiza automáticamente el estado en estados_programacion a 'FINALIZADO'.
    """
    semana = payload.get("semana")
    datos = payload.get("datos", {})
    usuario = payload.get("usuario", "Programacion")

    if not semana:
        return {"status": "error", "message": "Semana no especificada"}

    if not datos:
        return {"status": "error", "message": "No se recibieron datos de dotación para guardar."}

    try:
        exito, mensaje = modelo_dtprog.guardar_dotaciones(semana, datos)
        
        if exito:
            # Actualizamos el estado en la tabla estados_programacion
            modelo_dtprog.actualizar_estado_programacion(semana, estado="FINALIZADO", usuario=usuario)
            return {"status": "ok", "message": f"Dotación guardada correctamente para la Semana {semana}."}
        else:
            return {"status": "error", "message": mensaje}
    except Exception as e:
        print(f"Error en /api/dtprogramada/guardar: {e}")
        return {"status": "error", "message": str(e)}


@app.get("/api/dtprogramada/exportar-pdf")
@app.get("/generar_dt_pdf")
def api_exportar_pdf_dtprogramada(planta: Optional[str] = "ambas", semana: Optional[str] = None):
    """
    Genera y descarga el reporte PDF de la Dotación Programada discriminado por planta.
    """
    if not semana or semana.lower() == "none" or semana.strip() == "":
        semana = str(modelo_dtprog.obtener_numero_semana_personalizado())

    try:
        sectores_choco = modelo_dtprog.obtener_sectores_rotativos_por_planta("Chocolate")
        sectores_golo = modelo_dtprog.obtener_sectores_rotativos_por_planta("Golosinas")
        turnos = modelo_dtprog.obtener_turnos_por_tipo("ROTATIVO") or ["06:00 a 14:00", "14:00 a 22:00", "22:00 a 06:00"]
        
        dotaciones = modelo_dtprog.obtener_dotaciones_semana(semana)

        planta_norm = (planta or "ambas").strip().lower()
        
        if planta_norm in ["chocolates", "choco", "chocolate"]:
            sectores_filtrados = [("CHOCOLATES", sectores_choco, "Chocolate")]
            titulo_planta = "PLANTA CHOCOLATES"
        elif planta_norm in ["golosinas", "golo"]:
            sectores_filtrados = [("GOLOSINAS", sectores_golo, "Golosinas")]
            titulo_planta = "PLANTA GOLOSINAS"
        else:
            sectores_filtrados = [("CHOCOLATES", sectores_choco, "Chocolate"), ("GOLOSINAS", sectores_golo, "Golosinas")]
            titulo_planta = "CONSOLIDADO AMBAS PLANTAS"

        output = BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )
        elements = []
        styles = getSampleStyleSheet()

        elements.append(
            Paragraph(
                f"<b>DOTACIÓN PROGRAMADA POR TURNO - SEMANA {semana}</b>",
                styles["Title"],
            )
        )
        elements.append(
            Paragraph(
                f"<font size=11 color='#555555'><b>{titulo_planta}</b></font>",
                styles["Normal"],
            )
        )
        elements.append(Spacer(1, 14))

        for nombre_bloque, lista_sectores, clave_planta in sectores_filtrados:
            if not lista_sectores:
                continue

            elements.append(Paragraph(f"<b>{nombre_bloque}</b>", styles["Heading2"]))
            elements.append(Spacer(1, 6))

            encabezados = ["LÍNEA / SECTOR"] + [str(t) for t in turnos]
            tabla_datos = [encabezados]

            dot_planta = dotaciones.get(clave_planta, {})

            for s in lista_sectores:
                nombre_sec = s.get("nombre_sector") if isinstance(s, dict) else (s[0] if isinstance(s, (list, tuple)) else str(s))
                nombre_sec_clean = str(nombre_sec).strip()

                fila = [nombre_sec_clean]
                for t in turnos:
                    cant = dot_planta.get(nombre_sec_clean, {}).get(str(t).strip(), 0)
                    fila.append(str(cant) if cant is not None else "0")

                tabla_datos.append(fila)

            cant_turnos = len(turnos)
            ancho_sector = 200
            ancho_turno = (535 - ancho_sector) / cant_turnos if cant_turnos > 0 else 100
            col_widths = [ancho_sector] + [ancho_turno] * cant_turnos

            tabla_pdf = Table(tabla_datos, colWidths=col_widths)
            tabla_pdf.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003366")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("ALIGN", (0, 0), (0, -1), "LEFT"),
                    ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D3D3D3")),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                ])
            )

            elements.append(tabla_pdf)
            elements.append(Spacer(1, 16))

        doc.build(elements)
        output.seek(0)

        nombre_archivo = f"Dotacion_Programada_{planta_norm}_semana_{semana}.pdf"
        headers = {"Content-Disposition": f'inline; filename="{nombre_archivo}"'}
        return StreamingResponse(output, headers=headers, media_type="application/pdf")

    except Exception as e:
        print(f"Error al generar PDF de dotación programada: {e}")
        return {"error": str(e)}
# ==========================================
# 5.1. EXPORTACIÓN PDF: DOTACIÓN PROGRAMADA
# ==========================================

@app.get("/api/dtprogramada/exportar-pdf")
@app.get("/generar_dt_pdf")
def api_exportar_pdf_dtprogramada(planta: Optional[str] = "ambas", semana: Optional[str] = None):
    """
    Genera y descarga un reporte PDF con la Dotación Programada.
    Parámetros:
      - planta: 'chocolates', 'golosinas' o 'ambas' (por defecto)
      - semana: Número de semana (si se omite, calcula la semana activa)
    """
    if not semana or semana.lower() == "none" or semana.strip() == "":
        semana = semana_objetivo_por_defecto()

    try:
        # 1. Obtener datos desde el modelo de Dotación Programada
        sectores_choco = modelo_dtprog.obtener_sectores_rotativos_por_planta("Chocolate")
        sectores_golo = modelo_dtprog.obtener_sectores_rotativos_por_planta("Golosinas")
        turnos = modelo_dtprog.obtener_turnos_por_tipo("ROTATIVO") or ["06:00 a 14:00", "14:00 a 22:00", "22:00 a 06:00"]
        dotaciones = modelo_dtprog.obtener_dotaciones_semana(semana)

        # 2. Filtrar sectores según la planta solicitada
        planta_norm = (planta or "ambas").strip().lower()
        
        if planta_norm in ["chocolates", "choco", "chocolate"]:
            sectores_filtrados = [("CHOCOLATES", sectores_choco)]
            titulo_planta = "PLANTA CHOCOLATES"
        elif planta_norm in ["golosinas", "golo"]:
            sectores_filtrados = [("GOLOSINAS", sectores_golo)]
            titulo_planta = "PLANTA GOLOSINAS"
        else:
            sectores_filtrados = [("CHOCOLATES", sectores_choco), ("GOLOSINAS", sectores_golo)]
            titulo_planta = "CONSOLIDADO AMBAS PLANTAS"

        # Map de dotaciones guardadas por (sector, turno)
        map_dotacion = {}
        for reg in dotaciones:
            # Asumiendo estructura de tu modelo o tuplas
            if isinstance(reg, dict):
                sec, tur, cant = reg.get("sector"), reg.get("turno"), reg.get("cantidad", 0)
            elif len(reg) >= 3:
                sec, tur, cant = reg[0], reg[1], reg[2]
            else:
                continue
            map_dotacion[(str(sec).strip().upper(), str(tur).strip())] = cant

        # 3. Armar documento ReportLab
        output = BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )
        elements = []
        styles = getSampleStyleSheet()

        # Encabezado principal
        elements.append(
            Paragraph(
                f"<b>DOTACIÓN PROGRAMADA POR TURNO - SEMANA {semana}</b>",
                styles["Title"],
            )
        )
        elements.append(
            Paragraph(
                f"<font size=11 color='#555555'><b>{titulo_planta}</b></font>",
                styles["Normal"],
            )
        )
        elements.append(Spacer(1, 14))

        # Construcción de tablas por bloque de planta
        for nombre_bloque, lista_sectores in sectores_filtrados:
            if not lista_sectores:
                continue

            elements.append(Paragraph(f"<b>{nombre_bloque}</b>", styles["Heading2"]))
            elements.append(Spacer(1, 6))

            # Columnas: LÍNEA / SECTOR + Turnos
            encabezados = ["LÍNEA / SECTOR"] + [str(t) for t in turnos]
            tabla_datos = [encabezados]

            for s in lista_sectores:
                nombre_sec = s.get("nombre_sector") if isinstance(s, dict) else (s[0] if isinstance(s, (list, tuple)) else str(s))
                nombre_sec_clean = str(nombre_sec).strip()

                fila = [nombre_sec_clean]
                for t in turnos:
                    cant = map_dotacion.get((nombre_sec_clean.upper(), str(t).strip()), 0)
                    fila.append(str(cant) if cant is not None else "0")

                tabla_datos.append(fila)

            # Anchos de columna proporcionales
            cant_turnos = len(turnos)
            ancho_sector = 200
            ancho_turno = (535 - ancho_sector) / cant_turnos if cant_turnos > 0 else 100
            col_widths = [ancho_sector] + [ancho_turno] * cant_turnos

            tabla_pdf = Table(tabla_datos, colWidths=col_widths)
            tabla_pdf.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003366")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("ALIGN", (0, 0), (0, -1), "LEFT"),
                    ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D3D3D3")),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                ])
            )

            elements.append(tabla_pdf)
            elements.append(Spacer(1, 16))

        doc.build(elements)
        output.seek(0)

        nombre_archivo = f"Dotacion_Programada_{planta_norm}_semana_{semana}.pdf"
        headers = {"Content-Disposition": f'inline; filename="{nombre_archivo}"'}
        return StreamingResponse(output, headers=headers, media_type="application/pdf")

    except Exception as e:
        print(f"Error al generar PDF de dotación programada: {e}")
        return {"error": str(e)}


# ==========================================
# 6. MÓDULO: CHOCOLATES (SEMANAL Y FINDE)
# ==========================================

@app.get("/api/hchoco/datos")
def api_obtener_hchoco(semana: Optional[str] = None):
    modelo = ModeloHChocoWeb()
    modelo_base = modeloprincipal()
    try:
        if not semana or semana.lower() == "none" or semana.strip() == "":
            semana = str(modelo_base.obtener_numero_semana_personalizado(date.today()))

        modelo_base.eliminar_datos_semanas_anteriores(semana_actual=semana)

        sectores = modelo.obtener_sectores_chocolates()
        dotacion_teorica = modelo.obtener_dotacion_teorica(semana)
        datos_guardados = modelo.obtener_datos_semana(semana)
        operarios_disponibles = modelo.obtener_operarios_disponibles(semana)
        todos_operarios = modelo.obtener_operarios()

        sectores_info = []
        for row in sectores:
            if isinstance(row, dict):
                sec = row.get("nombre_sector") or row.get("sector")
                tipo = row.get("tipo_horario")
                req = row.get("requiere_dotacion")
            else:
                sec, tipo, req = row[0], row[1], row[2]

            t_horario = str(tipo).strip().upper() if tipo else "ROTATIVO"
            turnos = modelo.obtener_turnos_por_tipo(t_horario)
            if not turnos:
                turnos = ["08:00 a 16:00"] if t_horario == "CENTRAL" else ["06:00 a 14:00", "14:00 a 22:00", "22:00 a 06:00"]

            sectores_info.append({
                "sector": sec,
                "tipo_horario": t_horario,
                "requiere_dotacion": req,
                "turnos": turnos
            })

        return {
            "semana": semana,
            "sectores": sectores_info,
            "dotacion_teorica": dotacion_teorica,
            "asignados": datos_guardados,
            "operarios": todos_operarios,
            "operarios_disponibles": operarios_disponibles
        }
    except Exception as e:
        print(f"Error en endpoint /api/hchoco/datos: {e}")
        return {"error": str(e)}
    finally:
        modelo.cerrar_conexion()
        if hasattr(modelo_base, 'cerrar_conexion'):
            modelo_base.cerrar_conexion()


@app.post("/api/hchoco/guardar")
def api_guardar_hchoco(payload: dict = Body(...)):
    semana = payload.get("semana")
    registros = payload.get("registros", [])
    modelo = ModeloHChocoWeb()
    try:
        registros_tuplas = [(r[0], r[1], r[2], r[3]) for r in registros]
        modelo.guardar_datos_semana(semana, registros_tuplas)
        return {"status": "ok"}
    except Exception as e:
        print(f"Error en /api/hchoco/guardar: {e}")
        return {"status": "error", "mensaje": str(e)}
    finally:
        modelo.cerrar_conexion()


@app.get("/api/hchoco/finde/datos")
def api_obtener_hchoco_finde(semana: Optional[str] = None):
    modelo = ModeloHChocoFindeWeb()
    modelo_main = ModeloHChocoWeb()
    modelo_base = modeloprincipal()
    try:
        if not semana or semana.lower() == "none" or semana.strip() == "":
            semana = str(modelo_base.obtener_numero_semana_personalizado(date.today()))

        modelo_base.eliminar_datos_semanas_anteriores(semana_actual=semana)

        datos = modelo.obtener_datos_semana_finde(semana)
        operarios_disponibles = modelo.obtener_operarios_disponibles(semana)
        sectores = modelo_main.obtener_sectores_chocolates()

        return {
            "semana": semana,
            "asignados_finde": datos,
            "operarios_disponibles": operarios_disponibles,
            "sectores": sectores
        }
    except Exception as e:
        print(f"Error en /api/hchoco/finde/datos: {e}")
        return {"error": str(e)}
    finally:
        modelo.cerrar_conexion()
        modelo_main.cerrar_conexion()
        if hasattr(modelo_base, 'cerrar_conexion'):
            modelo_base.cerrar_conexion()


@app.post("/api/hchoco/finde/guardar")
def api_guardar_hchoco_finde(payload: dict = Body(...)):
    semana = payload.get("semana")
    registros = payload.get("registros", [])

    modelo = ModeloHChocoFindeWeb()
    try:
        modelo.guardar_datos_finde(semana, registros)
        return {"status": "ok"}
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"Error detallado en /api/hchoco/finde/guardar: {e}")
        return {"status": "error", "mensaje": str(e)}
    finally:
        modelo.cerrar_conexion()


@app.get("/api/hchoco/exportar-pdf")
def api_exportar_pdf_hchoco(semana: Optional[str] = None):
    if not semana:
        semana = semana_objetivo_por_defecto()

    modelo_main = ModeloHChocoWeb()
    modelo_finde = ModeloHChocoFindeWeb()

    try:
        output = BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )
        elements = []
        styles = getSampleStyleSheet()

        elements.append(
            Paragraph(
                f"<b>Semana {semana} - PLANTA CHOCOLATES (SEMANAL)</b>",
                styles["Title"],
            )
        )
        elements.append(Spacer(1, 12))

        datos_guardados = modelo_main.obtener_datos_semana(semana)
        sectores_operarios = {}

        for id_val, operario, horario, sector in datos_guardados:
            if sector and horario and operario:
                if sector not in sectores_operarios:
                    sectores_operarios[sector] = {}
                if horario not in sectores_operarios[sector]:
                    sectores_operarios[sector][horario] = []
                sectores_operarios[sector][horario].append(operario)

        for sector, turnos in sectores_operarios.items():
            if all(len(lst) == 0 for lst in turnos.values()):
                continue

            elements.append(Paragraph(f"<b>{sector}</b>", styles["Heading2"]))
            elements.append(Spacer(1, 6))

            columnas = list(turnos.keys())
            data = [columnas]
            max_filas = max(len(turnos[t]) for t in columnas) if columnas else 0

            for i in range(max_filas):
                fila = [
                    turnos[t][i] if i < len(turnos[t]) else "" for t in columnas
                ]
                data.append(fila)

            ancho_col = 535 / len(columnas) if columnas else 180
            table = Table(data, colWidths=[ancho_col] * len(columnas))
            table.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightblue),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ])
            )

            elements.append(table)
            elements.append(Spacer(1, 12))

        registros_finde = modelo_finde.obtener_datos_semana_finde(semana)
        if registros_finde:
            elements.append(PageBreak())
            elements.append(
                Paragraph(
                    f"<b>Semana {semana} - FIN DE SEMANA (CHOCOLATES)</b>",
                    styles["Title"],
                )
            )
            elements.append(Spacer(1, 12))

            data_finde = [["OPERARIO", "DÍA", "HORARIO", "SECTOR"]]
            for reg in registros_finde:
                if len(reg) == 5:
                    _, op, dia, hor, sec = reg
                elif len(reg) == 4:
                    op, dia, hor, sec = reg[1], "Fin de Semana", reg[2], reg[3]
                else:
                    continue

                if op:
                    data_finde.append([
                        str(op).strip(),
                        str(dia).strip(),
                        str(hor).strip(),
                        str(sec).strip(),
                    ])

            if len(data_finde) > 1:
                table_finde = Table(data_finde, colWidths=[180, 95, 110, 150])
                table_finde.setStyle(
                    TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAD3")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, 0), 10),
                        ("ALIGN", (0, 0), (0, -1), "LEFT"),
                        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 1), (-1, -1), 9),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ])
                )
                elements.append(table_finde)

        doc.build(elements)
        output.seek(0)

        nombre_archivo = f"Horarios chocolates semana {semana}.pdf"
        headers = {"Content-Disposition": f'attachment; filename="{nombre_archivo}"'}
        return StreamingResponse(output, headers=headers, media_type="application/pdf")
    finally:
        modelo_main.cerrar_conexion()
        modelo_finde.cerrar_conexion()


# ==========================================
# 7. MÓDULO: GOLOSINAS (SEMANAL Y FINDE)
# ==========================================

@app.get("/api/hgolo/datos")
def api_obtener_hgolo(semana: Optional[str] = None):
    modelo = ModeloHGoloWeb()
    modelo_base = modeloprincipal()
    try:
        if not semana or semana.lower() == "none" or semana.strip() == "":
            semana = str(modelo_base.obtener_numero_semana_personalizado(date.today()))

        modelo_base.eliminar_datos_semanas_anteriores(semana_actual=semana)

        sectores = modelo.obtener_sectores_golosinas()
        dotacion_teorica = modelo.obtener_dotacion_teorica(semana)
        datos_guardados = modelo.obtener_datos_semana(semana)
        operarios_disponibles = modelo.obtener_operarios_disponibles(semana)
        todos_operarios = modelo.obtener_operarios()

        sectores_info = []
        for row in sectores:
            if isinstance(row, dict):
                sec = row.get("nombre_sector") or row.get("sector")
                tipo = row.get("tipo_horario")
                req = row.get("requiere_dotacion")
            else:
                sec, tipo, req = row[0], row[1], row[2]

            t_horario = str(tipo).strip().upper() if tipo else "ROTATIVO"
            turnos = modelo.obtener_turnos_por_tipo(t_horario)
            if not turnos:
                turnos = ["08:00 a 16:00"] if t_horario == "CENTRAL" else ["06:00 a 14:00", "14:00 a 22:00", "22:00 a 06:00"]

            sectores_info.append({
                "sector": sec,
                "tipo_horario": t_horario,
                "requiere_dotacion": req,
                "turnos": turnos
            })

        return {
            "semana": semana,
            "sectores": sectores_info,
            "dotacion_teorica": dotacion_teorica,
            "asignados": datos_guardados,
            "operarios": todos_operarios,
            "operarios_disponibles": operarios_disponibles
        }
    except Exception as e:
        print(f"Error en endpoint /api/hgolo/datos: {e}")
        return {"error": str(e)}
    finally:
        modelo.cerrar_conexion()
        if hasattr(modelo_base, 'cerrar_conexion'):
            modelo_base.cerrar_conexion()


@app.post("/api/hgolo/guardar")
def api_guardar_hgolo(payload: dict = Body(...)):
    semana = payload.get("semana")
    registros = payload.get("registros", [])
    modelo = ModeloHGoloWeb()
    try:
        registros_tuplas = [(r[0], r[1], r[2], r[3]) for r in registros]
        modelo.guardar_datos_semana(semana, registros_tuplas)
        return {"status": "ok"}
    except Exception as e:
        print(f"Error en /api/hgolo/guardar: {e}")
        return {"status": "error", "mensaje": str(e)}
    finally:
        modelo.cerrar_conexion()


@app.get("/api/hgolo/finde/datos")
def api_hgolo_finde_datos(semana: Optional[str] = ""):
    modelo = ModeloHGoloFindeWeb()
    modelo_base = modeloprincipal()
    try:
        if not semana or semana.lower() == "none" or semana.strip() == "":
            semana = str(modelo_base.obtener_numero_semana_personalizado(date.today()))

        modelo_base.eliminar_datos_semanas_anteriores(semana_actual=semana)

        operarios = modelo.obtener_operarios_disponibles(semana)
        sectores = modelo.obtener_sectores()
        asignados = modelo.obtener_datos_semana_finde(semana)
        return {
            "semana": semana,
            "operarios_disponibles": operarios,
            "sectores": sectores,
            "asignados_finde": asignados
        }
    finally:
        modelo.cerrar_conexion()
        if hasattr(modelo_base, 'cerrar_conexion'):
            modelo_base.cerrar_conexion()


@app.post("/api/hgolo/finde/guardar")
async def api_hgolo_finde_guardar(request: Request):
    data = await request.json()
    semana = data.get("semana")
    registros = data.get("registros")

    modelo = ModeloHGoloFindeWeb()
    try:
        modelo.guardar_datos_finde(semana, registros)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}
    finally:
        modelo.cerrar_conexion()


@app.get("/api/hgolo/exportar-pdf")
def api_exportar_pdf_hgolo(semana: Optional[str] = None):
    if not semana:
        semana = semana_objetivo_por_defecto()

    modelo_main = ModeloHGoloWeb()
    modelo_finde = ModeloHGoloFindeWeb()

    try:
        output = BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )
        elements = []
        styles = getSampleStyleSheet()

        elements.append(
            Paragraph(
                f"<b>Semana {semana} - PLANTA GOLOSINAS (SEMANAL)</b>",
                styles["Title"],
            )
        )
        elements.append(Spacer(1, 12))

        datos_guardados = modelo_main.obtener_datos_semana(semana)
        sectores_operarios = {}

        for id_val, operario, horario, sector in datos_guardados:
            if sector and horario and operario:
                if sector not in sectores_operarios:
                    sectores_operarios[sector] = {}
                if horario not in sectores_operarios[sector]:
                    sectores_operarios[sector][horario] = []
                sectores_operarios[sector][horario].append(operario)

        for sector, turnos in sectores_operarios.items():
            if all(len(lst) == 0 for lst in turnos.values()):
                continue

            elements.append(Paragraph(f"<b>{sector}</b>", styles["Heading2"]))
            elements.append(Spacer(1, 6))

            columnas = list(turnos.keys())
            data = [columnas]
            max_filas = max(len(turnos[t]) for t in columnas) if columnas else 0

            for i in range(max_filas):
                fila = [
                    turnos[t][i] if i < len(turnos[t]) else "" for t in columnas
                ]
                data.append(fila)

            ancho_col = 535 / len(columnas) if columnas else 180
            table = Table(data, colWidths=[ancho_col] * len(columnas))
            table.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightblue),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ])
            )

            elements.append(table)
            elements.append(Spacer(1, 12))

        registros_finde = modelo_finde.obtener_datos_semana_finde(semana)
        if registros_finde:
            elements.append(PageBreak())
            elements.append(
                Paragraph(
                    f"<b>Semana {semana} - FIN DE SEMANA (GOLOSINAS)</b>",
                    styles["Title"],
                )
            )
            elements.append(Spacer(1, 12))

            data_finde = [["OPERARIO", "DÍA", "HORARIO", "SECTOR"]]
            for reg in registros_finde:
                if len(reg) == 5:
                    _, op, dia, hor, sec = reg
                elif len(reg) == 4:
                    op, dia, hor, sec = reg[1], "Fin de Semana", reg[2], reg[3]
                else:
                    continue

                if op:
                    data_finde.append([
                        str(op).strip(),
                        str(dia).strip(),
                        str(hor).strip(),
                        str(sec).strip(),
                    ])

            if len(data_finde) > 1:
                table_finde = Table(data_finde, colWidths=[180, 95, 110, 150])
                table_finde.setStyle(
                    TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAD3")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, 0), 10),
                        ("ALIGN", (0, 0), (0, -1), "LEFT"),
                        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 1), (-1, -1), 9),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ])
                )
                elements.append(table_finde)

        doc.build(elements)
        output.seek(0)

        nombre_archivo = f"Horarios golosinas semana {semana}.pdf"
        headers = {"Content-Disposition": f'attachment; filename="{nombre_archivo}"'}
        return StreamingResponse(output, headers=headers, media_type="application/pdf")
    finally:
        modelo_main.cerrar_conexion()
        modelo_finde.cerrar_conexion()


# ==========================================
# 8. MÓDULO: MECÁNICOS Y MANTENIMIENTO
# ==========================================

@app.get("/api/mecanicos/datos")
def api_mecanicos_datos(semana: Optional[str] = ""):
    modelo = ModeloMecanicosWeb()
    modelo_base = modeloprincipal()
    try:
        if not semana or semana.lower() == "none" or semana.strip() == "":
            semana = str(modelo_base.obtener_numero_semana_personalizado(date.today()))

        mecanicos_disponibles = modelo.obtener_mecanicos_disponibles()
        asignados = modelo.obtener_asignaciones_semana(semana)

        return {
            "semana": semana,
            "disponibles": mecanicos_disponibles,
            "asignados": asignados
        }
    finally:
        modelo.cerrar_conexion()
        if hasattr(modelo_base, 'cerrar_conexion'):
            modelo_base.cerrar_conexion()


@app.post("/api/mecanicos/guardar")
async def api_mecanicos_guardar(request: Request):
    data = await request.json()
    semana = data.get("semana")
    registros = data.get("registros", [])

    modelo = ModeloMecanicosWeb()
    try:
        modelo.guardar_asignaciones_mecanicos(semana, registros)
        return {"status": "ok"}
    except Exception as e:
        return {"status": "error", "message": str(e)}
    finally:
        modelo.cerrar_conexion()
